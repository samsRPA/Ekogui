import io
from typing import List
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from app.domain.interfaces.IExcelReader import IExcelReader


class OpenpyxlExcelReader(IExcelReader):
    def readFirstColumn(self, content: bytes) -> List[str]:
        try:
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        except (InvalidFileException, BadZipFile, KeyError, OSError) as e:
            raise ValueError(f"Could not read the file as an .xlsx workbook: {e}")

        try:
            if not workbook.worksheets:
                raise ValueError("The workbook has no sheets")

            sheet = workbook.worksheets[0]
            values: List[str] = []
            for (cell,) in sheet.iter_rows(min_col=1, max_col=1, values_only=True):
                if cell is None:
                    continue
                text = str(cell).strip()
                if text:
                    values.append(text)
            return values
        finally:
            workbook.close()
